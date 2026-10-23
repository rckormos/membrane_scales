import io
import json
from pathlib import Path
import zipfile

import numpy as np
import openpyxl
import pytest

from membrane_scales import BRC, MembraneScales, demo_pdb, normalize_sequence, pdb_coordinates, pdb_inventory

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "data/scales.json").read_text())
SCALES = MembraneScales(DATA)


@pytest.fixture(scope="module")
def workbook():
    with zipfile.ZipFile(ROOT / "sciadv.ads6804_data_s1.zip") as archive:
        return openpyxl.load_workbook(io.BytesIO(archive.read("ads6804_data_s1.xlsx")), data_only=True)


def test_all_sequence_coefficients_match_published_workbook(workbook):
    sheet = workbook["∆Gapp Per Residue"]
    for key, value_col, error_col in [("vitro_int", 2, 3), ("vitro_tm", 4, 5),
                                     ("vivo_int", 6, 7), ("vivo_tm", 8, 9),
                                     ("hwvh_7_5", 10, None), ("hwvh_4_5", 11, None),
                                     ("ww_octanol", 12, None), ("ww_bilayer", 13, None)]:
        for row in range(3, 23):
            aa = sheet.cell(row, 1).value
            assert DATA["sequence_scales"][key]["values"][aa] == sheet.cell(row, value_col).value
            if error_col:
                assert DATA["sequence_scales"][key]["stderr"][aa] == sheet.cell(row, error_col).value


def test_every_depth_value_and_reverse_boltzmann_reference(workbook):
    sheet = workbook["∆G PDB"]
    aa_order = "ARNDCQEGHILKMFPSTWYV"
    for population, first in [("all", 3), ("exposed", 46), ("buried", 89)]:
        for offset, aa in enumerate(aa_order):
            expected = np.array([sheet.cell(row, 51+offset).value for row in range(first, first+40)])
            np.testing.assert_array_equal(DATA["depth_scales"][population]["energy"][aa], expected)
            # Reconstruct Eq. 12 independently from exposed aqueous bins for every population.
            aqueous = np.mean([sheet.cell(row, 3+offset).value for row in [46, 47, 48, 83, 84, 85]])
            propensities = np.array([sheet.cell(row, 3+offset).value for row in range(first, first+40)])
            reconstructed = -(303 * 8.314 / 4184) * np.log(propensities / aqueous)
            np.testing.assert_allclose(expected, reconstructed, rtol=1e-12, atol=1e-12)
            for method in ["spline", "linear", "bin"]:
                result = SCALES.depth(aa*40, list(range(-39, 40, 2)), population, method)
                np.testing.assert_allclose(result.energy_kcal_mol, expected)


def test_published_triplet_examples_and_shared_coefficient_error():
    summary, _ = SCALES.sequence("DLP")
    scores = summary.set_index("scale")
    assert scores.loc["vitro_int", "sum_kcal_mol"] == pytest.approx(0.095056)
    assert scores.loc["vitro_tm", "sum_kcal_mol"] == pytest.approx(-0.150597)
    assert scores.loc["vitro_int_to_tm", "sum_kcal_mol"] == pytest.approx(-0.245653)
    triple, _ = SCALES.sequence("AAA")
    a = triple.set_index("scale").loc["vitro_int"]
    assert a.sum_kcal_mol == pytest.approx(-1.98168)
    assert a.approx_sum_se == pytest.approx(3 * 0.006989)
    probabilities = SCALES.scaffold_prediction("DLP")
    expected = np.exp(-np.array([0., 0.095056, -0.150597]) / 0.602)
    np.testing.assert_allclose(probabilities.model_probability, expected/expected.sum())
    np.testing.assert_allclose(SCALES.scaffold_prediction(BRC).model_probability, expected/expected.sum())


def test_missing_coefficients_never_become_partial_totals():
    summary, residues = SCALES.sequence("AMLQR")
    vivo = summary.set_index("scale").loc["vivo_int"]
    assert vivo.missing_aa == "A,M,Q,R"
    assert np.isnan(vivo.sum_kcal_mol)
    assert np.isnan(vivo.mean_kcal_mol)
    assert np.isnan(vivo.approx_sum_se)
    assert residues.vivo_int.isna().sum() == 4
    with pytest.raises(ValueError, match="lacks coefficients"):
        SCALES.scaffold_prediction("AAA", "vivo")
    with pytest.raises(ValueError, match="scaffold"):
        SCALES.scaffold_prediction("LLLALLL")


@pytest.mark.parametrize("text", ["", ">empty", ">one\nAAA\n>two\nCCC", "AA-X", "AA1", "BZJUO", "AAA*"])
def test_bad_sequence_rejected(text):
    with pytest.raises(ValueError):
        normalize_sequence(text)


def test_fasta_and_whitespace():
    assert normalize_sequence(">one helix\n awl \n dg\t") == "AWLDG"


def test_interpolation_edges_asymmetry_and_rt():
    energy = DATA["depth_scales"]["exposed"]["energy"]["W"]
    # -38 is exactly between the first two centers, but a boundary for binning.
    assert SCALES.depth("W", [-38], method="linear").energy_kcal_mol.iloc[0] == pytest.approx((energy[0]+energy[1])/2)
    assert SCALES.depth("W", [-38], method="bin").energy_kcal_mol.iloc[0] == energy[1]
    bins = SCALES.depth("WW", [-40, 40], method="bin")
    np.testing.assert_allclose(bins.energy_kcal_mol, [energy[0], energy[-1]])
    signed = SCALES.depth("WW", [-39, 39])
    assert signed.energy_kcal_mol.iloc[0] != signed.energy_kcal_mol.iloc[1]
    changed = SCALES.depth("WW", [-39, 39], rt="methods")
    np.testing.assert_allclose(changed.energy_kcal_mol, signed.energy_kcal_mol * 0.593 / (303*8.314/4184))
    with pytest.raises(ValueError, match="outside"):
        SCALES.depth("WW", [-39.01, 39.01])
    with pytest.warns(UserWarning, match="Clipped 2"):
        clipped = SCALES.depth("WW", [-100, 100], bounds="clip")
    assert clipped.clipped.all()
    np.testing.assert_allclose(clipped.energy_kcal_mol, [energy[0], energy[-1]])
    for bad_z in [[np.nan], [np.inf], [1, 2], [[1]]]:
        with pytest.raises(ValueError):
            SCALES.depth("W", bad_z)


def atom(serial, name, aa, chain, resid, insertion, xyz, alt=" ", occupancy=1.):
    x, y, z = xyz
    return (f"ATOM  {serial:5d} {name:^4s}{alt}{aa:3s} {chain}{resid:4d}{insertion}   "
            f"{x:8.3f}{y:8.3f}{z:8.3f}{occupancy:6.2f}{0.:6.2f}          C  ")


def fixture_pdb():
    return "\n".join([
        atom(1, "CA", "ALA", "A", 10, " ", (0, 0, -7)),
        atom(2, "CB", "ALA", "A", 10, " ", (0, 1, -6), "A", 0.4),
        atom(3, "CB", "ALA", "A", 10, " ", (0, 1, -8), "B", 0.6),
        atom(4, "CA", "GLY", "A", 10, "A", (3, 0, -5)),
        atom(5, "CA", "TRP", "A", 11, " ", (6, 0, -3)),
        atom(6, "CB", "TRP", "A", 11, " ", (6, 1, -2)),
        "TER   ", "END   ", ""])


def test_pdb_atom_conventions_insertions_alternate_atoms_and_scores():
    coordinates = pdb_coordinates(fixture_pdb())
    assert list(coordinates.aa) == ["A", "G", "W"]
    assert list(coordinates.residue_id) == ["10", "10A", "11"]
    assert list(coordinates.depth_atom) == ["CB", "CA", "CB"]
    assert list(coordinates.altloc) == ["B", "", ""]
    np.testing.assert_allclose(coordinates.z_cart_A, [-8, -5, -2])
    summary, rows = SCALES.score_pdb(coordinates, exposure_labels=["exposed", "buried", "all"])
    for population in ["exposed", "buried", "all"]:
        expected = SCALES.depth("AGW", [-8, -5, -2], population).energy_kcal_mol.to_numpy()
        np.testing.assert_allclose(rows[f"{population}_kcal_mol"], expected)
        assert summary.set_index("population").loc[population, "sum_kcal_mol"] == pytest.approx(expected.sum())
    np.testing.assert_allclose(rows.selected_kcal_mol, [rows.exposed_kcal_mol[0], rows.buried_kcal_mol[1], rows.all_kcal_mol[2]])
    selection = pdb_coordinates(fixture_pdb(), start="10A", end="11")
    assert "".join(selection.aa) == "GW"
    assert len(pdb_inventory(fixture_pdb())) == 1


def test_membrane_frame_rotation_and_translation_invariance():
    coordinates = pdb_coordinates(fixture_pdb())
    _, baseline = SCALES.score_pdb(coordinates)
    rotation = np.array([[0, 0, 1], [1, 0, 0], [0, 1, 0]], dtype=float)
    center = np.array([10., -21., 37.])
    moved = coordinates.copy()
    moved[["x_A", "y_A", "z_cart_A"]] = coordinates[["x_A", "y_A", "z_cart_A"]].to_numpy() @ rotation.T + center
    _, transformed = SCALES.score_pdb(moved, center=center, normal=3*(rotation @ [0, 0, 1]))
    np.testing.assert_allclose(transformed.z_A, baseline.z_A)
    np.testing.assert_allclose(transformed.exposed_kcal_mol, baseline.exposed_kcal_mol)
    with pytest.raises(ValueError, match="nonzero"):
        SCALES.score_pdb(coordinates, normal=(0, 0, 0))
    with pytest.raises(ValueError, match="finite"):
        SCALES.score_pdb(coordinates, center=(0, np.nan, 0))


def test_missing_atoms_unsupported_residues_and_selections_fail():
    no_cb = "\n".join(line for line in fixture_pdb().splitlines() if not (line[12:16].strip() == "CB" and line[22:26].strip() == "11"))
    with pytest.raises(ValueError, match="Missing CB"):
        pdb_coordinates(no_cb)
    no_gly_ca = "\n".join(line for line in fixture_pdb().splitlines() if line[17:20] != "GLY")
    with pytest.raises(ValueError, match="gap"):
        pdb_coordinates(no_gly_ca)
    with pytest.raises(ValueError, match="Unsupported"):
        pdb_coordinates(fixture_pdb().replace("TRP", "MSE"))
    with pytest.raises(ValueError, match="endpoint"):
        pdb_coordinates(fixture_pdb(), start="9")
    with pytest.raises(ValueError, match="after"):
        pdb_coordinates(fixture_pdb(), start="11", end="10")
    with pytest.raises(ValueError, match="Chain"):
        pdb_coordinates(fixture_pdb(), chain_id="Z")
    with pytest.raises(ValueError, match="model_index"):
        pdb_coordinates(fixture_pdb(), model_index=-1)


def test_model_selection_and_demo():
    atoms = "\n".join(line for line in fixture_pdb().splitlines() if line.startswith("ATOM"))
    models = f"MODEL       10\n{atoms}\nENDMDL\nMODEL       20\n{atoms}\nENDMDL\nEND   \n"
    assert pdb_inventory(models).pdb_model_serial.tolist() == [10, 20]
    assert pdb_coordinates(models, model_index=1).model_index.tolist() == [1, 1, 1]
    coordinates = pdb_coordinates(demo_pdb())
    assert "".join(coordinates.aa) == BRC
    assert set(coordinates.depth_atom) == {"CB"}
    summary, _ = SCALES.score_pdb(coordinates)
    assert np.isfinite(summary.sum_kcal_mol).all()


def test_fitted_backend_requires_author_model_and_validates_output():
    with pytest.raises(ValueError, match="not been supplied"):
        SCALES.depth("A", [0], method="fitted")
    calls = []
    def fake_model(population, aa, z):
        calls.append((population, aa))
        return z * 0.1 + (1 if aa == "A" else 2)
    adapter = MembraneScales(DATA, fitted_model=fake_model)
    result = adapter.depth("AGA", [-5, 0, 5], method="fitted")
    np.testing.assert_allclose(result.energy_kcal_mol, [0.5, 2, 1.5])
    assert set(calls) == {("exposed", "A"), ("exposed", "G")}
    invalid = MembraneScales(DATA, fitted_model=lambda population, aa, z: float("nan"))
    with pytest.raises(ValueError, match="finite array"):
        invalid.depth("A", [0], method="fitted")


def test_natural_spline_against_analytic_solution():
    data = json.loads(json.dumps(DATA))
    # Natural spline through (-1,0), (0,1), (1,0) is
    # 1 - 1.5*z**2 + 0.5*abs(z)**3, with zero endpoint curvature.
    for table in data["depth_scales"].values():
        table["z"] = [-1., 0., 1.]
        table["energy"] = {aa: [0., 1., 0.] for aa in table["energy"]}
    scales = MembraneScales(data)
    z = np.linspace(-1, 1, 21)
    expected = 1 - 1.5*z**2 + 0.5*np.abs(z)**3
    for population in ["exposed", "buried", "all"]:
        result = scales.depth("AGW"*7, z, population)
        np.testing.assert_allclose(result.energy_kcal_mol, expected, atol=1e-14)
    assert scales.depth("A", [0.5]).energy_kcal_mol.iloc[0] == pytest.approx(0.6875)
    assert scales.depth("A", [0.5], method="linear").energy_kcal_mol.iloc[0] == 0.5
    for outside in [-1.01, 1.01]:
        with pytest.raises(ValueError, match="outside"):
            scales.depth("A", [outside])
    with pytest.warns(UserWarning, match="Clipped"):
        clipped = scales.depth("AA", [-2, 2], bounds="clip")
    np.testing.assert_allclose(clipped.energy_kcal_mol, [0., 0.], atol=1e-14)
    assert clipped.clipped.all()


@pytest.mark.parametrize("selection, method", [
    ("Natural cubic splines", "spline"),
    ("Piecewise linear interpolation", "linear"),
    ("Gaussian/sigmoid fits (to be implemented)", "fitted"),
])
def test_notebook_dropdown_routes_scoring_and_rejects_pending_fit(selection, method):
    notebook = json.loads((ROOT / "membrane_scales_colab.ipynb").read_text())
    source = next("".join(cell["source"]) for cell in notebook["cells"]
                  if "#@title PDB selection and scoring parameters" in "".join(cell["source"]))
    dropdown = next(line for line in source.splitlines() if line.startswith("interpolation_method ="))
    assert json.loads(dropdown.split("#@param", 1)[1]) == [
        "Natural cubic splines", "Piecewise linear interpolation", "Gaussian/sigmoid fits (to be implemented)"]
    source = source.replace('interpolation_method = "Natural cubic splines"', f'interpolation_method = {selection!r}', 1)
    ns = {"pdb_source": "demo", "pdb_text": demo_pdb(), "scales": SCALES,
          "pdb_coordinates": pdb_coordinates, "display": lambda value: None,
          "include_comparison_scales": True}
    if method == "fitted":
        with pytest.raises(ValueError, match="not yet implemented"):
            exec(source, ns)
        assert ns["pdb_summary"] is None
    else:
        exec(source, ns)
        assert ns["depth_method"] == method
        assert ns["pdb_summary"].method.tolist() == [method]*3
        expected = SCALES.depth(BRC, ns["pdb_residues"].z_A, method=method)
        np.testing.assert_allclose(ns["pdb_residues"].exposed_kcal_mol, expected.energy_kcal_mol)


def test_notebook_embeds_exact_tested_source_and_data():
    notebook = json.loads((ROOT / "membrane_scales_colab.ipynb").read_text())
    sources = ["".join(cell["source"]) for cell in notebook["cells"]]
    module = (ROOT / "membrane_scales.py").read_text().strip()
    assert any(source.endswith(module) for source in sources)
    data_cell = next(source for source in sources if "DATA = json.loads(" in source)
    namespace = {"json": json, "MembraneScales": MembraneScales}
    exec(data_cell, namespace)
    assert namespace["DATA"] == DATA
