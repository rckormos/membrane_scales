"""Extract the published, full-precision coefficients without refitting them."""
from pathlib import Path
import hashlib
import io
import json
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
AA3 = dict(zip(
    "ALA ARG ASN ASP CYS GLN GLU GLY HIS ILE LEU LYS MET PHE PRO SER THR TRP TYR VAL".split(),
    "ARNDCQEGHILKMFPSTWYV",
))


def extract():
    archive = ROOT / "sciadv.ads6804_data_s1.zip"
    with zipfile.ZipFile(archive) as zipped:
        raw = zipped.read("ads6804_data_s1.xlsx")
    workbook = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    sheet = workbook["∆Gapp Per Residue"]
    definitions = [
        ("vitro_int", "In vitro: Wat → Int", 2, 3, "experimental"),
        ("vitro_tm", "In vitro: Wat → TM", 4, 5, "experimental"),
        ("vivo_int", "In vivo: Wat → Int", 6, 7, "experimental"),
        ("vivo_tm", "In vivo: Wat → TM", 8, 9, "experimental"),
        ("hwvh_7_5", "HWvH at 7.5 Å (Data S1 comparison)", 10, None, "comparison"),
        ("hwvh_4_5", "HWvH at 4.5 Å (Data S1 comparison)", 11, None, "comparison"),
        ("ww_octanol", "Wimley–White octanol (as tabulated in Data S1)", 12, None, "comparison"),
        ("ww_bilayer", "Wimley–White bilayer (as tabulated in Data S1)", 13, None, "comparison"),
    ]
    scales = {}
    for key, label, col, error_col, kind in definitions:
        scales[key] = {
            "label": label, "kind": kind,
            "source": f"∆Gapp Per Residue!{openpyxl.utils.get_column_letter(col)}3:{openpyxl.utils.get_column_letter(col)}22",
            "values": {sheet.cell(r, 1).value: sheet.cell(r, col).value for r in range(3, 23)},
            "stderr": ({sheet.cell(r, 1).value: sheet.cell(r, error_col).value for r in range(3, 23)}
                       if error_col else None),
        }
    sheet = workbook["∆G PDB"]
    depth = {}
    for kind, first in [("all", 3), ("exposed", 46), ("buried", 89)]:
        depth[kind] = {
            "source": f"∆G PDB!AW{first}:BR{first+39}",
            "z": [sheet.cell(r, 49).value for r in range(first, first+40)],
            "energy": {AA3[sheet.cell(first-1, c).value]:
                       [sheet.cell(r, c).value for r in range(first, first+40)]
                       for c in range(51, 71)},
            "propensity": {AA3[sheet.cell(first-1, c).value]:
                           [sheet.cell(r, c).value for r in range(first, first+40)]
                           for c in range(3, 23)},
        }
    result = {
        "citation": "Grau et al. Science Advances 11, eads6804 (2025)",
        "doi": "10.1126/sciadv.ads6804",
        "workbook_sha256": hashlib.sha256(raw).hexdigest(),
        "archive_sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
        "units": "kcal/mol",
        "pdb_rt_workbook": 303 * 8.314 / 4184,
        "pdb_rt_methods": 0.593,
        "experimental_rt": {"vitro": 0.602, "vivo": 0.612},
        "sequence_scales": scales,
        "depth_scales": depth,
        "notes": [
            "Data S1 is authoritative for the numerical coefficients in this implementation.",
            "Supplementary Table 1 rounds coefficients and appears to interchange E/P Wat→TM entries relative to Data S1.",
            "The workbook uses RT=303*8.314/4184; main-text Eq. 12 states RT=0.593 kcal/mol.",
            "Fig. 3 mentions sigmoid/Gaussian fits but fit coefficients are absent from the supplied files. No original fitted model is reconstructed.",
            "In vivo coefficients for A, M, Q, R are absent and remain null.",
        ],
    }
    destination = ROOT / "data" / "scales.json"
    destination.parent.mkdir(exist_ok=True)
    destination.write_text(json.dumps(result, indent=2, ensure_ascii=False, allow_nan=False) + "\n")
    print(destination)


if __name__ == "__main__":
    extract()
